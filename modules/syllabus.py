"""Syllabus text normalization and bounded local document extraction."""

from __future__ import annotations

import io
import re
import unicodedata
import zipfile
from pathlib import PurePath


MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_PDF_PAGES = 500
MAX_DOCX_ENTRIES = 1000
MAX_DOCX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024


class SyllabusError(ValueError):
    """Raised for invalid, unsupported, or unsafe syllabus input."""


def normalize_syllabus_text(text: str) -> str:
    """Normalize spacing while preserving only words present in source text."""
    if not isinstance(text, str):
        raise SyllabusError("Syllabus text must be a string.")
    normalized = unicodedata.normalize("NFC", text).replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t\f\v]+", " ", line).strip() for line in normalized.split("\n")]
    while lines and not lines[0]:
        lines.pop(0)
    while lines and not lines[-1]:
        lines.pop()
    return "\n".join(lines)


def parse_syllabus_text(text: str) -> dict[str, object]:
    """Normalize text into source-preserving lines usable as syllabus topics.

    Each non-empty source line becomes one topic; no topic is inferred or added.
    """
    normalized = normalize_syllabus_text(text)
    topics = [line for line in normalized.splitlines() if line]
    return {"source_text": normalized, "topics": topics}


def extract_syllabus_text(
    filename: str,
    content: bytes,
    max_bytes: int = MAX_UPLOAD_BYTES,
) -> str:
    """Extract text from bounded TXT, PDF, or DOCX bytes without storing uploads.

    PDF extraction uses ``pypdf``; DOCX extraction uses ``python-docx``.
    These packages are imported only when their format is requested.
    """
    if not isinstance(filename, str) or not isinstance(content, bytes):
        raise SyllabusError("A filename and uploaded file bytes are required.")
    if isinstance(max_bytes, bool) or not isinstance(max_bytes, int) or max_bytes <= 0:
        raise SyllabusError("max_bytes must be a positive integer.")
    if not content:
        raise SyllabusError("The uploaded file is empty.")
    if len(content) > min(max_bytes, MAX_UPLOAD_BYTES):
        raise SyllabusError(f"Upload exceeds the {min(max_bytes, MAX_UPLOAD_BYTES)} byte limit.")
    extension = PurePath(filename).suffix.lower()
    if extension not in {".txt", ".pdf", ".docx"}:
        raise SyllabusError("Only .txt, .pdf, and .docx syllabus files are supported.")

    if extension == ".txt":
        try:
            extracted = normalize_syllabus_text(content.decode("utf-8-sig"))
        except UnicodeDecodeError as error:
            raise SyllabusError("TXT syllabus files must use UTF-8 encoding.") from error
    elif extension == ".pdf":
        extracted = _extract_pdf(content)
    else:
        extracted = _extract_docx(content)
    if not extracted:
        raise SyllabusError("No readable text was found in the syllabus file.")
    return extracted


def _extract_pdf(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as error:
        raise SyllabusError("PDF extraction requires the 'pypdf' package.") from error
    try:
        reader = PdfReader(io.BytesIO(content), strict=True)
        if reader.is_encrypted:
            raise SyllabusError("Encrypted PDF files are not supported.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise SyllabusError(
                f"PDF syllabi may contain no more than {MAX_PDF_PAGES} pages."
            )
        pages = [page.extract_text() or "" for page in reader.pages]
    except SyllabusError:
        raise
    except Exception as error:
        raise SyllabusError(f"Could not extract text from PDF: {error}") from error
    return normalize_syllabus_text("\n".join(pages))


def _extract_docx(content: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_DOCX_ENTRIES:
                raise SyllabusError("DOCX contains too many archive entries.")
            if sum(entry.file_size for entry in entries) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise SyllabusError("DOCX expands beyond the allowed extraction limit.")
            if any(entry.flag_bits & 0x1 for entry in entries):
                raise SyllabusError("Encrypted DOCX files are not supported.")
            if "word/document.xml" not in archive.namelist():
                raise SyllabusError("The uploaded DOCX does not contain a document.")
    except SyllabusError:
        raise
    except (OSError, zipfile.BadZipFile, RuntimeError) as error:
        raise SyllabusError("The uploaded DOCX is invalid.") from error
    try:
        from docx import Document
    except ImportError as error:
        raise SyllabusError("DOCX extraction requires the 'python-docx' package.") from error
    try:
        document = Document(io.BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            paragraphs.extend(
                cell.text for row in table.rows for cell in row.cells
            )
    except Exception as error:
        raise SyllabusError(f"Could not extract text from DOCX: {error}") from error
    return normalize_syllabus_text("\n".join(paragraphs))
